"""Planilhas prontas por tema (XLSX), sem inventar dados.

"Jarvis, crie uma planilha de gastos" -> Documentos\\Jarvis\\Planilhas\\
gastos-20260918-1630.xlsx com colunas, formatos (R$, datas, %), totais por
formula, cabecalho congelado e filtros. Linhas de exemplo so quando pedido
("com exemplo"), numa aba separada e rotulada como ficticia. Nunca
sobrescreve um arquivo existente.

Os 11 temas cobrem os de referencia (tarefas, estoque, alunos, clientes,
funcionarios, gastos, receitas, projeto, agenda, saude, generico); colunas e
formulas foram escritas aqui.
"""

from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

MOEDA = 'R$ #,##0.00'
DATA = 'DD/MM/YYYY'
HORA = 'HH:MM'
PCT = '0%'

# tema -> (titulo, [(coluna, largura, formato)], totais: {coluna: "soma"|"media"}, formulas por linha)
MODELOS: dict[str, dict[str, Any]] = {
    "tarefas": {"titulo": "Tarefas", "colunas": [("Tarefa", 36, None), ("Responsável", 18, None), ("Prioridade", 12, None),
                ("Prazo", 13, DATA), ("Status", 14, None), ("Observações", 36, None)],
                "listas": {"Prioridade": ["Alta", "Média", "Baixa"], "Status": ["A fazer", "Fazendo", "Feito"]}},
    "estoque": {"titulo": "Estoque", "colunas": [("Item", 30, None), ("Código", 12, None), ("Categoria", 16, None),
                ("Quantidade", 12, '0'), ("Mínimo", 10, '0'), ("Custo unitário", 15, MOEDA), ("Valor em estoque", 17, MOEDA),
                ("Repor?", 9, None), ("Local", 16, None)],
                "linha": {"Valor em estoque": "=D{r}*F{r}", "Repor?": '=IF(D{r}="","",IF(D{r}<E{r},"SIM",""))'},
                "totais": {"Valor em estoque": "soma"}},
    "alunos": {"titulo": "Notas dos alunos", "colunas": [("Aluno", 30, None), ("Turma", 10, None), ("Nota 1", 9, '0.0'),
               ("Nota 2", 9, '0.0'), ("Nota 3", 9, '0.0'), ("Média", 9, '0.0'), ("Situação", 13, None)],
               "linha": {"Média": '=IF(COUNT(C{r}:E{r})=0,"",AVERAGE(C{r}:E{r}))',
                         "Situação": '=IF(F{r}="","",IF(F{r}>=6,"Aprovado","Recuperação"))'},
               "totais": {"Média": "media"}},
    "clientes": {"titulo": "Clientes", "colunas": [("Nome", 30, None), ("Telefone", 16, '@'), ("E-mail", 28, None),
                 ("Cidade", 18, None), ("Cliente desde", 14, DATA), ("Observações", 36, None)]},
    "funcionarios": {"titulo": "Funcionários", "colunas": [("Nome", 30, None), ("Cargo", 20, None), ("Setor", 16, None),
                     ("Admissão", 13, DATA), ("Salário", 14, MOEDA), ("Contato", 20, '@')],
                     "totais": {"Salário": "soma"}},
    "gastos": {"titulo": "Gastos", "colunas": [("Data", 13, DATA), ("Descrição", 34, None), ("Categoria", 16, None),
               ("Forma de pagamento", 20, None), ("Valor", 14, MOEDA)],
               "listas": {"Categoria": ["Moradia", "Alimentação", "Transporte", "Saúde", "Lazer", "Educação", "Outros"],
                          "Forma de pagamento": ["Pix", "Débito", "Crédito", "Dinheiro", "Boleto"]},
               "totais": {"Valor": "soma"}},
    "receitas": {"titulo": "Receitas", "colunas": [("Data", 13, DATA), ("Fonte", 22, None), ("Descrição", 34, None),
                 ("Valor", 14, MOEDA)], "totais": {"Valor": "soma"}},
    "projeto": {"titulo": "Projeto", "colunas": [("Etapa", 32, None), ("Responsável", 18, None), ("Início", 13, DATA),
                ("Fim", 13, DATA), ("Dias", 7, '0'), ("% concluído", 12, PCT), ("Status", 14, None)],
                "linha": {"Dias": '=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r})'},
                "listas": {"Status": ["Não iniciado", "Em andamento", "Concluído", "Atrasado"]}},
    "agenda": {"titulo": "Agenda", "colunas": [("Data", 13, DATA), ("Hora", 8, HORA), ("Compromisso", 34, None),
               ("Local", 20, None), ("Com quem", 20, None), ("Observações", 30, None)]},
    "saude": {"titulo": "Saúde", "colunas": [("Data", 13, DATA), ("Peso (kg)", 11, '0.0'), ("Pressão", 11, None),
              ("Sono (h)", 10, '0.0'), ("Atividade (min)", 15, '0'), ("Água (L)", 10, '0.0'), ("Observações", 30, None)],
              "totais": {"Peso (kg)": "media", "Sono (h)": "media", "Atividade (min)": "soma"}},
    "generico": {"titulo": "Planilha", "colunas": [("Item", 30, None), ("Descrição", 36, None), ("Quantidade", 12, '0'),
                 ("Valor", 14, MOEDA), ("Observações", 30, None)], "totais": {"Valor": "soma"}},
}

_PALAVRAS = {
    "tarefas": r"tarefas?|afazeres|to ?do|pendencias|atividades",
    "estoque": r"estoque|inventario|produtos|mercadorias",
    "alunos": r"alunos?|notas|turma|escola|boletim",
    "clientes": r"clientes?|contatos|crm",
    "funcionarios": r"funcionarios?|colaboradores?|equipe|folha",
    "gastos": r"gastos?|despesas?|financas|financeira|orcamento|contas a pagar|custos",
    "receitas": r"receitas?|renda|ganhos|entradas|faturamento|vendas",
    "projeto": r"projetos?|cronograma|etapas",
    "agenda": r"agenda|compromissos|eventos|reunioes",
    "saude": r"saude|peso|treino|academia|exercicios?|sono|pressao",
}


def _norm(t: str) -> str:
    t = unicodedata.normalize("NFKD", t.lower())
    return "".join(c for c in t if not unicodedata.combining(c))


def escolher_modelo(pedido: str) -> str:
    n = _norm(pedido)
    for tema, pal in _PALAVRAS.items():
        if re.search(rf"\b(?:{pal})\b", n):
            return tema
    return "generico"


def pasta_documentos() -> Path:
    """Pasta Documentos de verdade (pode estar redirecionada ao OneDrive)."""
    try:
        import ctypes
        import uuid

        class GUID(ctypes.Structure):
            _fields_ = [("d1", ctypes.c_ulong), ("d2", ctypes.c_ushort), ("d3", ctypes.c_ushort), ("d4", ctypes.c_ubyte * 8)]
        g = uuid.UUID("{FDD39AD0-238F-46AF-ADB4-6C85480369C7}")
        guid = GUID(g.fields[0], g.fields[1], g.fields[2], (ctypes.c_ubyte * 8)(*g.bytes[8:]))
        p = ctypes.c_wchar_p()
        if ctypes.windll.shell32.SHGetKnownFolderPath(ctypes.byref(guid), 0, None, ctypes.byref(p)) == 0:
            caminho = Path(p.value)
            ctypes.windll.ole32.CoTaskMemFree(p)
            return caminho
    except (OSError, AttributeError, ValueError):
        pass
    return Path.home() / "Documents"


def caminho_livre(pasta: Path, base: str, ext: str) -> Path:
    """Nunca sobrescreve: acrescenta -2, -3... se ja existir."""
    pasta.mkdir(parents=True, exist_ok=True)
    alvo = pasta / f"{base}{ext}"
    n = 2
    while alvo.exists():
        alvo = pasta / f"{base}-{n}{ext}"
        n += 1
    return alvo


def _exemplo(tema: str, colunas: list[tuple[str, int, str | None]]) -> list[list[Any]]:
    """Tres linhas claramente ficticias (nomes "Exemplo A/B/C")."""
    hoje = date.today()
    linhas = []
    for i, rot in enumerate("ABC"):
        linha = []
        for nome, _, fmt in colunas:
            if fmt == DATA:
                linha.append(hoje + timedelta(days=i))
            elif fmt == HORA:
                linha.append(f"{9 + i}:00")
            elif fmt == MOEDA:
                linha.append(round(50.0 * (i + 1), 2))
            elif fmt in ("0", "0.0"):
                linha.append(i + 1 if fmt == "0" else 6.0 + i)
            elif fmt == PCT:
                linha.append(0.25 * (i + 1))
            else:
                linha.append(f"Exemplo {rot} ({nome.lower()})")
        linhas.append(linha)
    return linhas


def criar(pedido: str, pasta: Path | None = None, com_exemplo: bool | None = None) -> dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    tema = escolher_modelo(pedido)
    m = MODELOS[tema]
    if com_exemplo is None:
        com_exemplo = bool(re.search(r"\bcom (?:dados de )?exemplos?\b|\bpreenchid\w*\b", _norm(pedido)))
    pasta = pasta or (pasta_documentos() / "Jarvis" / "Planilhas")
    arquivo = caminho_livre(pasta, f"{tema}-{datetime.now():%Y%m%d-%H%M}", ".xlsx")

    wb = Workbook()
    abas = [(wb.active, m["titulo"][:31], False)]
    if com_exemplo:
        abas.append((wb.create_sheet(), "Exemplo (fictício)", True))
    nomes = [c[0] for c in m["colunas"]]
    linhas_uteis = 200
    for ws, titulo, exemplo in abas:
        ws.title = titulo
        ws.append(nomes)
        for i, (nome, largura, fmt) in enumerate(m["colunas"], start=1):
            letra = get_column_letter(i)
            ws.column_dimensions[letra].width = largura
            c = ws.cell(row=1, column=i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center", vertical="center")
        dados = _exemplo(tema, m["colunas"]) if exemplo else []
        ultima = 1 + (len(dados) if exemplo else linhas_uteis)
        for r in range(2, ultima + 1):
            for i, (nome, _, fmt) in enumerate(m["colunas"], start=1):
                c = ws.cell(row=r, column=i)
                if exemplo and nome not in m.get("linha", {}):
                    c.value = dados[r - 2][i - 1]
                if nome in m.get("linha", {}):
                    c.value = m["linha"][nome].format(r=r)
                if fmt:
                    c.number_format = fmt
        for nome, opcoes in m.get("listas", {}).items():
            letra = get_column_letter(nomes.index(nome) + 1)
            dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letra}2:{letra}{ultima}")
        if m.get("totais"):
            rt = ultima + 1
            ws.cell(row=rt, column=1, value="Total" if any(v == "soma" for v in m["totais"].values()) else "Média").font = Font(bold=True)
            for nome, tipo in m["totais"].items():
                i = nomes.index(nome) + 1
                letra = get_column_letter(i)
                f = "SUM" if tipo == "soma" else "AVERAGE"
                c = ws.cell(row=rt, column=i, value=f'=IFERROR({f}({letra}2:{letra}{ultima}),"")')
                c.font = Font(bold=True)
                c.number_format = m["colunas"][i - 1][2] or "General"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(nomes))}{ultima}"
        if exemplo:
            ws.sheet_properties.tabColor = "C00000"
            ws.cell(row=ultima + 3, column=1, value="Dados FICTÍCIOS, só para mostrar o formato. Apague esta aba quando quiser.").font = Font(italic=True, color="C00000")
    wb.save(arquivo)
    return {"tema": tema, "titulo": m["titulo"], "arquivo": str(arquivo), "colunas": nomes, "exemplo": com_exemplo}


def fala(r: dict[str, Any]) -> str:
    ex = " Pus uma aba de exemplo com dados fictícios." if r["exemplo"] else ""
    cols = ", ".join(r["colunas"][:4]) + ("…" if len(r["colunas"]) > 4 else "")
    return f"Planilha de {r['titulo'].lower()} criada em Documentos, pasta Jarvis, Planilhas. Colunas: {cols}.{ex} Abrindo."
