"""Assistente de engenharia (como o Jarvis na oficina do Tony): contas de
eletrica e fisica, materiais, projeto de pecas (via cad) e analise de dados.

Tudo deterministico e local: as contas nao vao ao modelo (que erra conta).
Valores de materiais sao TIPICOS de referencia (variam com liga e fabricacao)
e a resposta diz isso.
"""

from __future__ import annotations

import csv
import math
import re
import statistics
import unicodedata
from pathlib import Path
from typing import Any

G = 9.81
NUM = r"(\d+(?:[.,]\d+)?)"


def _norm(t: str) -> str:
    t = unicodedata.normalize("NFKD", t.lower())
    return " ".join("".join(c for c in t if not unicodedata.combining(c)).split())


def _n(s: str) -> float:
    return float(s.replace(",", "."))


def fmt(v: float, casas: int = 2) -> str:
    """1234.5 -> "1.234,5" (sem zeros sobrando)."""
    if abs(v) >= 100:
        casas = min(casas, 1)
    s = f"{v:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if "," in s:
        s = s.rstrip("0").rstrip(",")
    return s


def preparar(frase: str) -> str:
    """Numeros por extenso -> digitos e texto sem acento (reusa a das contas)."""
    from .utilidades import digitos
    return digitos(frase)


# ---- eletrica -----------------------------------------------------------------------------
def eletrica(frase: str) -> str | None:
    """Lei de Ohm e potencia: com dois de tensao/corrente/resistencia/potencia, acha o pedido."""
    n = preparar(frase)
    pede = ("corrente" if re.search(r"\bcorrente\b|\bquantos? (?:amperes?|amps?)\b", n) else
            "resistencia" if re.search(r"\bresist\w*\b|\bquantos? ohms?\b|\bque resistor\b", n) else
            "tensao" if re.search(r"\btensao\b|\bvoltagem\b|\bquantos? volts?\b", n) else
            "potencia" if re.search(r"\bpotencia\b|\bquantos? (?:watts?|w)\b", n) else None)
    if not pede:
        return None
    v = re.search(NUM + r" ?(?:v|volts?)\b", n)
    i = re.search(NUM + r" ?(ma|miliamperes?|a|amperes?|amps?)\b", n)
    r = re.search(NUM + r" ?(k|kilo|quilo)? ?(?:ohms?|Ω)", n)
    p = re.search(NUM + r" ?(kw|quilowatts?|w|watts?)\b", n)
    V = _n(v.group(1)) if v else None
    I = (_n(i.group(1)) / (1000 if i.group(2).startswith("m") else 1)) if i else None
    R = (_n(r.group(1)) * (1000 if r.group(2) else 1)) if r else None
    P = (_n(p.group(1)) * (1000 if p.group(2).startswith(("k", "q")) else 1)) if p else None
    try:
        if pede == "corrente":
            res = V / R if V is not None and R else P / V if P is not None and V else math.sqrt(P / R) if P and R else None
            return None if res is None else f"A corrente é de {fmt(res, 3)} ampères ({fmt(res * 1000, 1)} miliampères)."
        if pede == "resistencia":
            res = V / I if V is not None and I else V * V / P if V and P else P / (I * I) if P and I else None
            return None if res is None else f"A resistência é de {fmt(res, 2)} ohms."
        if pede == "tensao":
            res = I * R if I is not None and R is not None else P / I if P and I else math.sqrt(P * R) if P and R else None
            return None if res is None else f"A tensão é de {fmt(res, 2)} volts."
        res = V * I if V is not None and I is not None else V * V / R if V and R else I * I * R if I and R else None
        return None if res is None else f"A potência é de {fmt(res, 2)} watts."
    except ZeroDivisionError:
        return "Com esses valores a conta não fecha (divisão por zero)."


def consumo(frase: str, tarifa: float | None = None) -> str | None:
    """ "quanto gasta um aparelho de 1500 watts ligado 2 horas por dia" -> kWh e (com tarifa) reais."""
    n = preparar(frase)
    if not re.search(r"\b(?:quanto (?:gasta|consome|custa)|consumo)\b", n):
        return None
    p = re.search(NUM + r" ?(kw|quilowatts?|w|watts?)\b", n)
    h = re.search(NUM + r" ?(?:h|horas?)\b", n)
    m = re.search(NUM + r" ?(?:min|minutos?)\b", n)
    if not p or not (h or m):
        return None
    watts = _n(p.group(1)) * (1000 if p.group(2).startswith(("k", "q")) else 1)
    horas = (_n(h.group(1)) if h else 0) + (_n(m.group(1)) / 60 if m else 0)
    por_dia = bool(re.search(r"\b(?:por|ao|todo) dia\b|\bdiari\w*\b", n))
    kwh = watts * horas / 1000
    if por_dia:
        mes = kwh * 30
        s = f"Dá {fmt(kwh, 2)} quilowatts-hora por dia, cerca de {fmt(mes, 1)} por mês"
        custo = mes * tarifa if tarifa else None
    else:
        s = f"Dá {fmt(kwh, 3)} quilowatts-hora"
        custo = kwh * tarifa if tarifa else None
    if custo is not None:
        return s + f", uns {fmt(custo, 2)} reais pela tarifa de {fmt(tarifa, 2)} por kWh."
    return s + ". Para eu calcular em reais, diga: a tarifa de luz é, por exemplo, 0,85."


def tarifa_dita(frase: str) -> float | None:
    m = re.search(r"\btarifa (?:de luz |de energia |do kwh |da luz )?(?:e|eh|esta|custa|agora e) (?:de )?(?:r\$ ?)?"
                  + NUM, preparar(frase))
    return _n(m.group(1)) if m else None


# ---- fisica -------------------------------------------------------------------------------------
def fisica(frase: str) -> tuple[str, dict[str, Any] | None] | None:
    """Queda livre, lancamento obliquo e energia cinetica. Devolve (fala, grafico opcional)."""
    n = preparar(frase)
    if re.search(r"\blanc\w*\b|\bchut\w*\b|\barremess\w*\b|\balcance\b", n):
        v = re.search(NUM + r" ?(?:m/s|metros? por segundo)", n)
        a = re.search(NUM + r" ?(?:graus|°)", n)
        if v and a:
            vel, ang = _n(v.group(1)), math.radians(_n(a.group(1)))
            alcance = vel * vel * math.sin(2 * ang) / G
            altura = (vel * math.sin(ang)) ** 2 / (2 * G)
            tempo = 2 * vel * math.sin(ang) / G
            pontos = [[vel * math.cos(ang) * t, vel * math.sin(ang) * t - G * t * t / 2]
                      for t in (tempo * k / 40 for k in range(41))]
            fala = (f"Sem resistência do ar: alcance de {fmt(alcance, 1)} metros, altura máxima de {fmt(altura, 1)} "
                    f"metros e {fmt(tempo, 2)} segundos no ar.")
            return fala, {"titulo": f"Trajetória · {fmt(vel, 1)} m/s a {fmt(_n(a.group(1)), 0)}° (metros)", "unidade": "m",
                          "series": [{"nome": "altura", "pontos": pontos}]}
    if re.search(r"\b(?:cai\w*|queda|solt\w*|larg\w*)\b", n):
        h = re.search(NUM + r" ?(?:m|metros?)\b(?! por)", n)
        if h:
            alt = _n(h.group(1))
            t = math.sqrt(2 * alt / G)
            return (f"Sem resistência do ar, leva {fmt(t, 2)} segundos para cair {fmt(alt, 1)} metros e chega a "
                    f"{fmt(G * t, 1)} metros por segundo, uns {fmt(G * t * 3.6, 0)} quilômetros por hora."), None
    if re.search(r"\benergia cinetica\b", n):
        m = re.search(NUM + r" ?(?:kg|quilos?|quilogramas?)\b", n)
        v = re.search(NUM + r" ?(?:m/s|metros? por segundo)", n)
        if m and v:
            e = _n(m.group(1)) * _n(v.group(1)) ** 2 / 2
            return f"A energia cinética é de {fmt(e, 1)} joules.", None
    return None


# ---- materiais ------------------------------------------------------------------------------------
# nome, apelidos, densidade g/cm3, tracao (MPa, texto), fusao/uso (texto)
MATERIAIS = [
    ("aço carbono", ("aco", "aco carbono", "aco comum"), 7.85, "400 a 550", "funde entre 1.425 e 1.540 °C"),
    ("aço inox 304", ("inox", "aco inox", "aco inoxidavel"), 8.0, "cerca de 515", "funde entre 1.400 e 1.450 °C"),
    ("alumínio 6061", ("aluminio",), 2.70, "cerca de 310", "funde entre 582 e 652 °C"),
    ("cobre", ("cobre",), 8.96, "cerca de 210", "funde a 1.085 °C"),
    ("latão", ("latao",), 8.5, "cerca de 340", "funde entre 900 e 940 °C"),
    ("bronze", ("bronze",), 8.8, "cerca de 300", "funde por volta de 950 °C"),
    ("titânio Ti-6Al-4V", ("titanio",), 4.43, "cerca de 950", "funde entre 1.604 e 1.660 °C"),
    ("ferro fundido", ("ferro fundido", "ferro"), 7.2, "cerca de 200", "funde entre 1.150 e 1.200 °C"),
    ("chumbo", ("chumbo",), 11.34, "cerca de 18", "funde a 327 °C"),
    ("ouro", ("ouro",), 19.32, "cerca de 120", "funde a 1.064 °C"),
    ("prata", ("prata",), 10.49, "cerca de 170", "funde a 962 °C"),
    ("PLA", ("pla",), 1.24, "50 a 60", "amolece por volta de 60 °C; imprime entre 190 e 220 °C"),
    ("ABS", ("abs",), 1.04, "cerca de 40", "amolece por volta de 105 °C; imprime entre 220 e 250 °C"),
    ("PETG", ("petg",), 1.27, "cerca de 50", "amolece por volta de 80 °C; imprime entre 220 e 250 °C"),
    ("nylon PA6", ("nylon", "nailon", "nilon"), 1.14, "70 a 80", "funde por volta de 220 °C"),
    ("vidro comum", ("vidro",), 2.5, "30 a 90, e quebra sem deformar", "amolece por volta de 700 °C"),
    ("concreto", ("concreto",), 2.4, "só 2 a 5 na tração (20 a 40 na compressão)", "não funde; perde resistência acima de 300 °C"),
    ("borracha natural", ("borracha",), 0.93, "20 a 30", "degrada acima de 200 °C"),
    ("fibra de carbono com epóxi", ("fibra de carbono", "carbono"), 1.6, "600 a 1.500, conforme a direção das fibras",
     "a resina amolece por volta de 120 °C"),
]


def materiais_citados(frase: str) -> list[tuple]:
    """Materiais na ordem em que aparecem; o apelido mais longo vence ("aco inox" nao e "aco")."""
    n = f" {_norm(frase)} "
    trechos = []
    for m in MATERIAIS:
        for apelido in m[1]:
            for achado in re.finditer(rf"(?<=\s){re.escape(apelido)}(?=\s)", n):
                trechos.append((achado.start(), achado.end(), m))
    aceitos: list[tuple[int, int, tuple]] = []
    for ini_, fim_, m in sorted(trechos, key=lambda t: t[0] - t[1]):          # mais longos primeiro
        if any(not (fim_ <= a0 or ini_ >= a1) for a0, a1, _ in aceitos) or any(m is x for _, _, x in aceitos):
            continue
        aceitos.append((ini_, fim_, m))
    return [m for _, _, m in sorted(aceitos, key=lambda t: t[0])]


def material(frase: str) -> str | None:
    n = preparar(frase)
    mats = materiais_citados(n)
    if not mats:
        return None
    if re.search(r"\bcompar\w*\b|\bdiferenca\b|\bqual (?:e )?(?:melhor|mais (?:leve|forte|resistente|pesado))\b", n) and len(mats) >= 2:
        a, b = mats[0], mats[1]
        leve = a if a[2] < b[2] else b
        partes = [f"{a[0]}: {fmt(a[2])} gramas por centímetro cúbico, tração de {a[3]} megapascais, {a[4]}",
                  f"{b[0]}: {fmt(b[2])} gramas por centímetro cúbico, tração de {b[3]} megapascais, {b[4]}"]
        razao = max(a[2], b[2]) / min(a[2], b[2])
        peso = (f"Pesam quase o mesmo" if razao < 1.15 else f"O {leve[0]} é {fmt(razao, 1)} vezes mais leve")
        return ". ".join(partes) + f". {peso}. São valores típicos: variam com a liga e a fabricação."
    m = mats[0]
    if re.search(r"\bdensidade\b|\bquanto pesa um (?:centimetro|metro) cubico\b", n):
        return f"A densidade típica do {m[0]} é {fmt(m[2])} gramas por centímetro cúbico."
    if re.search(r"\b(?:fusao|derret\w*|funde|temperatura)\b", n):
        return f"O {m[0]} {m[4]}."
    if re.search(r"\b(?:resistencia|aguenta|tracao|forte)\b", n):
        return f"A resistência à tração típica do {m[0]} é de {m[3]} megapascais."
    return None


def peso_de_peca(frase: str) -> str | None:
    """ "quanto pesa um cubo de alumínio de 10 cm" -> volume x densidade."""
    n = preparar(frase)
    if not re.search(r"\bquanto pesa\b|\bqual (?:e )?o peso\b|\bqual (?:e )?a massa\b", n):
        return None
    mats = materiais_citados(n)
    if not mats:
        return None
    mat = mats[0]
    def cm(valor: str, unidade: str | None) -> float:
        u = (unidade or "cm").strip()
        return _n(valor) * (0.1 if u.startswith(("mm", "mili")) else 100 if u in ("m", "metro", "metros") else 1)
    U = r" ?(mm|milimetros?|cm|centimetros?|m|metros?)?\b"
    vol = None
    if (m := re.search(r"\bcubo\b.*?" + NUM + U, n)):
        a = cm(m.group(1), m.group(2)); vol = a ** 3
    elif (m := re.search(r"\besfera\b.*?" + NUM + U, n)):
        d = cm(m.group(1), m.group(2)); vol = math.pi * d ** 3 / 6
    elif (m := re.search(r"\bcilindro\b", n)):
        d = re.search(NUM + U + r" de diametro|diametro de " + NUM + U, n)
        h = re.search(NUM + U + r" de (?:altura|comprimento)|(?:altura|comprimento) de " + NUM + U, n)
        if d and h:
            dv = cm(d.group(1) or d.group(3), d.group(2) or d.group(4)); hv = cm(h.group(1) or h.group(3), h.group(2) or h.group(4))
            vol = math.pi * dv * dv / 4 * hv
    elif (m := re.search(NUM + U + r" (?:por|x) " + NUM + U + r" (?:por|x) " + NUM + U, n)):
        vol = cm(m.group(1), m.group(2)) * cm(m.group(3), m.group(4)) * cm(m.group(5), m.group(6))
    if vol is None:
        return None
    massa = vol * mat[2]
    peso = f"{fmt(massa / 1000, 2)} quilos" if massa >= 1000 else f"{fmt(massa, 1)} gramas"
    return f"Com {fmt(vol, 1)} centímetros cúbicos de {mat[0]}, pesa cerca de {peso}."


# ---- pecas (CAD) ------------------------------------------------------------------------------------
def pedido_de_peca(frase: str) -> tuple[str, dict[str, float]] | str | None:
    """ "projete uma engrenagem de 20 dentes com 40 mm" -> ("engrenagem", {...}); str = pergunta de volta."""
    n = preparar(frase)
    U = r" ?(mm|milimetros?|cm|centimetros?)?"
    def mm(v: str, u: str | None) -> float:
        return _n(v) * (10 if u and u.startswith(("cm", "cent")) else 1)
    medidas = [mm(m.group(1), m.group(2)) for m in re.finditer(r"(?<![\d.,])" + NUM + r"(?![\d.,])" + U, n)
               if not re.match(r"\s*dentes", n[m.end():])]
    if re.search(r"\bengrenage\w*\b", n):
        d = re.search(r"(\d+) dentes", n)
        if not d:
            return "Quantos dentes, Senhor? Por exemplo: engrenagem de 20 dentes com 40 milímetros."
        dentes = int(d.group(1))
        esp = re.search(NUM + U + r" de (?:espessura|grossura|altura)|(?:espessura|grossura|altura) de " + NUM + U, n)
        furo = re.search(r"furo (?:de )?" + NUM + U, n)
        resto = medidas
        diam = next((x for x in resto if not (esp and x == mm(esp.group(1) or esp.group(3), esp.group(2) or esp.group(4)))
                     and not (furo and x == mm(furo.group(1), furo.group(2)))), None)
        if diam is None:
            return "Com qual diâmetro externo, Senhor? Por exemplo: 40 milímetros."
        return "engrenagem", {"dentes": dentes, "externo": diam,
                              "espessura": mm(esp.group(1) or esp.group(3), esp.group(2) or esp.group(4)) if esp else max(3.0, diam / 5),
                              "furo": mm(furo.group(1), furo.group(2)) if furo else round(diam / 8, 1)}
    if re.search(r"\bcaixa\b", n):
        m = re.search(NUM + U + r" (?:por|x) " + NUM + U + r" (?:por|x) " + NUM + U, n)
        if not m:
            return "Quais as medidas da caixa? Por exemplo: 50 por 30 por 20 milímetros."
        un = m.group(6) or m.group(4) or m.group(2)
        c, l, a = (mm(m.group(i), m.group(i + 1) or un) for i in (1, 3, 5))
        par = re.search(r"parede (?:de )?" + NUM + U, n)
        parede = mm(par.group(1), par.group(2)) if par else (2.0 if re.search(r"\b(?:oca|aberta|vazada|organizador\w*)\b", n) else None)
        return "caixa", {"c": c, "l": l, "a": a, "parede": parede or 0}
    if re.search(r"\btubo\b", n):
        if len(medidas) < 3:
            return "Diga o diâmetro de fora, o de dentro e a altura. Por exemplo: tubo de 30 por 26 milímetros com 40 de altura."
        return "tubo", {"externo": max(medidas[0], medidas[1]), "interno": min(medidas[0], medidas[1]), "altura": medidas[2]}
    if re.search(r"\b(?:cilindro|pino|disco)\b", n):
        d = re.search(NUM + U + r" de diametro|diametro de " + NUM + U, n)
        h = re.search(NUM + U + r" de (?:altura|comprimento|espessura)|(?:altura|comprimento|espessura) de " + NUM + U, n)
        if d and h:
            return "cilindro", {"diametro": mm(d.group(1) or d.group(3), d.group(2) or d.group(4)),
                                "altura": mm(h.group(1) or h.group(3), h.group(2) or h.group(4))}
        if len(medidas) >= 2:
            return "cilindro", {"diametro": medidas[0], "altura": medidas[1]}
        return "Diga o diâmetro e a altura. Por exemplo: cilindro de 20 milímetros de diâmetro e 50 de altura."
    if re.search(r"\b(?:esfera|bola)\b", n):
        return ("esfera", {"diametro": medidas[0]}) if medidas else "Qual o diâmetro da esfera?"
    if re.search(r"\bcone\b", n):
        return ("cone", {"diametro": medidas[0], "altura": medidas[1]}) if len(medidas) >= 2 else "Diga o diâmetro e a altura do cone."
    return None


def gerar_peca(tipo: str, p: dict[str, float]):
    from . import cad
    if tipo == "engrenagem":
        tris = cad.engrenagem(int(p["dentes"]), p["externo"], p["espessura"], p["furo"])
        nome = f"engrenagem {int(p['dentes'])} dentes {fmt(p['externo'], 1)} mm"
        desc = (f"engrenagem de {int(p['dentes'])} dentes, {fmt(p['externo'], 1)} milímetros de diâmetro, "
                f"{fmt(p['espessura'], 1)} de espessura e furo de {fmt(p['furo'], 1)}, módulo {fmt(p['externo'] / (p['dentes'] + 2), 2)}")
    elif tipo == "caixa":
        tris = cad.caixa(p["c"], p["l"], p["a"], p["parede"] or None)
        nome = f"caixa {fmt(p['c'], 0)}x{fmt(p['l'], 0)}x{fmt(p['a'], 0)} mm"
        desc = (f"caixa de {fmt(p['c'], 1)} por {fmt(p['l'], 1)} por {fmt(p['a'], 1)} milímetros"
                + (f", aberta, com parede de {fmt(p['parede'], 1)}" if p["parede"] else ", maciça"))
    elif tipo == "tubo":
        tris = cad.tubo(p["externo"], p["interno"], p["altura"])
        nome = f"tubo {fmt(p['externo'], 0)}-{fmt(p['interno'], 0)} x {fmt(p['altura'], 0)} mm"
        desc = f"tubo de {fmt(p['externo'], 1)} por fora, {fmt(p['interno'], 1)} por dentro e {fmt(p['altura'], 1)} de altura"
    elif tipo == "cilindro":
        tris = cad.cilindro(p["diametro"], p["altura"])
        nome = f"cilindro {fmt(p['diametro'], 0)} x {fmt(p['altura'], 0)} mm"
        desc = f"cilindro de {fmt(p['diametro'], 1)} milímetros de diâmetro e {fmt(p['altura'], 1)} de altura"
    elif tipo == "esfera":
        tris = cad.esfera(p["diametro"])
        nome = f"esfera {fmt(p['diametro'], 0)} mm"
        desc = f"esfera de {fmt(p['diametro'], 1)} milímetros"
    else:
        tris = cad.cone(p["diametro"], p["altura"])
        nome = f"cone {fmt(p['diametro'], 0)} x {fmt(p['altura'], 0)} mm"
        desc = f"cone de {fmt(p['diametro'], 1)} milímetros de base e {fmt(p['altura'], 1)} de altura"
    return tris, nome, desc


def falar_peca(desc: str, volume_mm3: float, caminho: Path) -> str:
    cm3 = volume_mm3 / 1000
    return (f"Pronto: {desc}. Volume de {fmt(cm3, 1)} centímetros cúbicos, uns {fmt(cm3 * 1.24, 1)} gramas em PLA "
            f"maciço. Salvei em Documentos, Jarvis, Projetos, e está na mesa holográfica. Para imprimir, diga: abra no Cura.")


# ---- dados ------------------------------------------------------------------------------------------
def ler_tabela(caminho: Path, limite: int = 20000) -> tuple[list[str], list[list[Any]]]:
    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(caminho, read_only=True, data_only=True)
        linhas = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True, max_row=limite + 1)]
        wb.close()
    else:
        bruto = caminho.read_bytes()[:20_000_000]
        texto = next((bruto.decode(c) for c in ("utf-8-sig", "cp1252") if _decodifica(bruto, c)), bruto.decode("latin-1"))
        dialeto = csv.Sniffer().sniff(texto[:4000], delimiters=",;\t") if texto.strip() else csv.excel
        linhas = [r for r in csv.reader(texto.splitlines(), dialeto)][: limite + 1]
    linhas = [l for l in linhas if any(c not in (None, "") for c in l)]
    if not linhas:
        return [], []
    cab = [str(c).strip() if c not in (None, "") else f"coluna {i + 1}" for i, c in enumerate(linhas[0])]
    return cab, linhas[1:]


def _decodifica(b: bytes, cod: str) -> bool:
    try:
        b.decode(cod)
        return True
    except UnicodeDecodeError:
        return False


def _numero(v: Any) -> float | None:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace("R$", "").replace("%", "").strip()
        if re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+(?:,\d+)?|-?\d+(?:,\d+)?", s):
            return float(s.replace(".", "").replace(",", "."))
        if re.fullmatch(r"-?\d+(?:\.\d+)?", s):
            return float(s)
    return None


def analisar_tabela(cab: list[str], linhas: list[list[Any]]) -> dict[str, Any]:
    """Colunas numericas: n, soma, media, min, max, desvio e tendencia (primeira vs ultima metade)."""
    colunas = []
    for i, nome in enumerate(cab):
        vals = [x for x in (_numero(l[i]) if i < len(l) else None for l in linhas) if x is not None]
        if len(vals) < max(2, 0.6 * len(linhas)):
            continue
        meio = len(vals) // 2
        tend = (statistics.fmean(vals[meio:]) / statistics.fmean(vals[:meio]) - 1) * 100 if meio and statistics.fmean(vals[:meio]) else 0
        colunas.append({"nome": nome, "n": len(vals), "soma": sum(vals), "media": statistics.fmean(vals),
                        "min": min(vals), "max": max(vals), "desvio": statistics.pstdev(vals), "tendencia": tend,
                        "valores": vals})
    return {"linhas": len(linhas), "colunas": colunas, "cabecalho": cab}


def falar_analise(nome_arquivo: str, a: dict[str, Any]) -> str:
    if not a["colunas"]:
        return f"{nome_arquivo} tem {a['linhas']} linhas, mas nenhuma coluna de números para analisar."
    partes = [f"{nome_arquivo}: {a['linhas']} linhas e {len(a['colunas'])} coluna{'s' if len(a['colunas']) > 1 else ''} de números."]
    for c in a["colunas"][:3]:
        t = c["tendencia"]
        tendencia = ("estável" if abs(t) < 3 else f"subindo {fmt(t, 0)} por cento" if t > 0 else f"caindo {fmt(-t, 0)} por cento")
        partes.append(f"{c['nome']}: soma {fmt(c['soma'])}, média {fmt(c['media'])}, de {fmt(c['min'])} a {fmt(c['max'])}, "
                      f"{tendencia} da primeira para a segunda metade.")
    return " ".join(partes)


def grafico_da_analise(nome_arquivo: str, a: dict[str, Any]) -> dict[str, Any] | None:
    if not a["colunas"]:
        return None
    c = a["colunas"][0]
    return {"titulo": f"{nome_arquivo} · {c['nome']}", "unidade": "",
            "series": [{"nome": c["nome"], "pontos": [[i, v] for i, v in enumerate(c["valores"][:500])]}]}
