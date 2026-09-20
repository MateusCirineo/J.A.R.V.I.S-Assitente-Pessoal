"""Planilhas por tema (F1) e programas gerados sem executar (F2)."""

import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
sys.path.append(str(RAIZ / "_libs"))

from openpyxl import load_workbook  # noqa: E402

from hud_runtime import comandos as cmd, planilhas, programas  # noqa: E402


class TestPlanilhas(unittest.TestCase):
    def test_escolhe_tema(self):
        casos = {"crie uma planilha de gastos do mês": "gastos", "planilha de controle de estoque": "estoque",
                 "planilha com as notas dos alunos": "alunos", "planilha de tarefas": "tarefas",
                 "planilha de clientes": "clientes", "planilha de funcionários": "funcionarios",
                 "planilha de receitas e vendas": "receitas", "cronograma do projeto": "projeto",
                 "planilha dos meus compromissos": "agenda", "planilha de treino e peso": "saude",
                 "uma planilha qualquer": "generico"}
        for frase, tema in casos.items():
            self.assertEqual(planilhas.escolher_modelo(frase), tema, frase)
        self.assertEqual(len(planilhas.MODELOS), 11)

    def test_sem_dados_inventados_e_com_formulas(self):
        pasta = Path(tempfile.mkdtemp())
        r = planilhas.criar("crie uma planilha de estoque", pasta)
        wb = load_workbook(r["arquivo"])
        self.assertEqual(wb.sheetnames, ["Estoque"])                 # nenhuma aba de exemplo
        ws = wb.active
        self.assertEqual(ws["A1"].value, "Item")
        self.assertIsNone(ws["A2"].value)                             # nada inventado
        self.assertEqual(ws["G2"].value, "=D2*F2")                    # valor em estoque
        self.assertIn("SUM(G2:G201)", ws["G202"].value)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws["F2"].number_format, planilhas.MOEDA)

    def test_exemplo_so_quando_pedido_e_rotulado(self):
        r = planilhas.criar("planilha de gastos com exemplo", Path(tempfile.mkdtemp()))
        wb = load_workbook(r["arquivo"])
        self.assertEqual(wb.sheetnames, ["Gastos", "Exemplo (fictício)"])
        ex = wb["Exemplo (fictício)"]
        self.assertTrue(str(ex["B2"].value).startswith("Exemplo A"))
        self.assertIsNone(wb["Gastos"]["B2"].value)
        self.assertIn("fictício", planilhas.fala(r))

    def test_nunca_sobrescreve(self):
        pasta = Path(tempfile.mkdtemp())
        a = planilhas.caminho_livre(pasta, "x", ".xlsx")
        a.write_text("original")
        b = planilhas.caminho_livre(pasta, "x", ".xlsx")
        self.assertNotEqual(a, b)
        self.assertEqual(b.name, "x-2.xlsx")
        self.assertEqual(a.read_text(), "original")


CODIGO_OK = "```python\nimport os\n\ndef main():\n    print('oi')\n\nif __name__ == '__main__':\n    main()\n```"
CODIGO_PERIGOSO = ("```python\nimport shutil, subprocess\n\ndef main():\n    shutil.rmtree('x')\n"
                   "    subprocess.run(['dir'])\n    eval('1+1')\n```")


class TestProgramas(unittest.TestCase):
    def test_valida_sem_executar(self):
        erro, avisos = programas.validar("import os\nos.remove('arquivo')\n")
        self.assertIsNone(erro)
        self.assertEqual(avisos, ["apaga arquivos"])
        erro, _ = programas.validar("def f(:\n  pass")
        self.assertIn("erro de sintaxe na linha 1", erro)

    def test_cria_abre_para_revisar_e_avisa(self):
        pasta = Path(tempfile.mkdtemp())
        abertos = []
        p = programas.Programas(lambda m: CODIGO_PERIGOSO, pasta, abrir=lambda a: abertos.append(a) or "VS Code")
        with mock.patch("os.startfile", create=True) as sf, mock.patch("subprocess.run") as run:
            r = p.criar("crie um programa em Python que limpa a pasta temporária")
            sf.assert_not_called()                                   # nunca abre com o interpretador
            run.assert_not_called()
        arq = Path(r["arquivo"])
        self.assertEqual(abertos, [arq])
        self.assertEqual(arq.parent, pasta)
        self.assertEqual(arq.suffix, ".py")
        texto = arq.read_text(encoding="utf-8")
        self.assertIn("# NÃO foi executado", texto)
        self.assertEqual(set(r["avisos"]), {"apaga arquivos", "roda comandos do sistema", "executa texto como código"})
        self.assertIn("Não executei nada", programas.fala(r))

    def test_segunda_chance_quando_sintaxe_falha(self):
        respostas = iter(["```python\ndef main(:\n```", CODIGO_OK])
        pedidos = []
        p = programas.Programas(lambda m: pedidos.append(m) or next(respostas), Path(tempfile.mkdtemp()),
                                abrir=lambda a: "Bloco de Notas")
        r = p.criar("escreva um script que diz oi")
        self.assertIsNone(r["erro"])
        self.assertIn("erro de sintaxe", pedidos[1][-1]["content"])

    def test_rotas(self):
        casos = {"Jarvis, crie uma planilha de gastos": "planilha", "planilha de clientes": "planilha",
                 "crie um programa em Python que renomeia fotos": "programa",
                 "escreva um script que baixa a cotação": "programa",
                 "abra o programa calculadora": "app"}
        for frase, esperado in casos.items():
            achado = cmd.interpretar(frase)
            self.assertEqual(achado[0] if achado else None, esperado, frase)


if __name__ == "__main__":
    unittest.main()
