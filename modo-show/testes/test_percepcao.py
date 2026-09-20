"""Percepcao sob pedido: janela (S1), documentos (S2), audio/video (S4)."""

import os
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
sys.path.append(str(RAIZ / "_libs"))

from hud_runtime import comandos as cmd, documentos, midia_arquivos, tela  # noqa: E402
from hud_runtime.visao import Visao  # noqa: E402

def _pdf(texto: bytes) -> bytes:
    """PDF minimo valido (com xref) contendo uma linha de texto."""
    conteudo = b"BT /F1 18 Tf 20 60 Td (" + texto + b") Tj ET"
    objs = [b"<</Type/Catalog/Pages 2 0 R>>", b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
            b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 144]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
            b"<</Length %d>>stream\n" % len(conteudo) + conteudo + b"\nendstream",
            b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>"]
    saida, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(saida))
        saida += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(saida)
    saida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    saida += b"".join(b"%010d 00000 n \n" % o for o in offs)
    saida += b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return saida


PDF = _pdf(b"Contrato de aluguel 2026")


class TestRotas(unittest.TestCase):
    def test_rotas(self):
        casos = {"Jarvis, resuma o último PDF baixado": "documento", "resuma o documento aberto": "documento",
                 "o que diz o contrato de aluguel?": "documento", "leia o arquivo notas": "documento",
                 r"resuma C:\x\a.pdf": "documento",
                 "transcreva o último áudio baixado": "midia", "analise o vídeo da reunião": "midia",
                 "pause o vídeo": "musica_pausar", "me lembre de ler o documento às 15h": "lembrete",
                 "Jarvis, olhe minha tela": "tela", "explique esse erro": "tela", "ligue a tela cheia": "tela_cheia",
                 "o que é isso?": "visao"}
        for frase, esperado in casos.items():
            achado = cmd.interpretar(frase)
            self.assertEqual(achado[0] if achado else None, esperado, frase)


class TestTela(unittest.TestCase):
    def test_janelas_do_jarvis_e_sensiveis(self):
        for t in ("J.A.R.V.I.S.", "Painel", "Jarvis runtime - feche para desligar", "JARVIS · Central de Comando"):
            self.assertTrue(tela.e_do_jarvis(t), t)
        for t in ("main.py - Jarvis - Visual Studio Code", "Documento1 - Word"):
            self.assertFalse(tela.e_do_jarvis(t), t)
        for t in ("Bitwarden", "Nubank - Microsoft Edge", "Gerenciador de Credenciais"):
            self.assertTrue(tela.sensivel(t), t)
        self.assertFalse(tela.sensivel("relatorio.xlsx - Excel"))

    def test_recusa_janela_de_banco_sem_capturar(self):
        estado = SimpleNamespace(atualizar=mock.Mock(), registrar=mock.Mock())
        v = Visao(estado, camera=None, prefs=SimpleNamespace(ler=lambda: {"modelo_voz": "m"}))
        with mock.patch.object(tela, "janela_de_trabalho", return_value=(1, "Itaú - Internet Banking")), \
             mock.patch.object(tela, "capturar") as cap:
            self.assertIn("não analiso janelas de senha ou de banco", v.analisar_tela())
            cap.assert_not_called()


class TestDocumentos(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp())

    def test_extrair_txt_docx_pdf(self):
        t = self.pasta / "notas.txt"
        t.write_bytes("Reunião às 15h".encode("cp1252"))
        self.assertEqual(documentos.extrair_texto(t)[0], "Reunião às 15h")
        import docx
        d = docx.Document()
        d.add_paragraph("Relatório de vendas")
        d.add_paragraph("Total: R$ 10 mil")
        d.save(self.pasta / "relatorio.docx")
        self.assertIn("Total: R$ 10 mil", documentos.extrair_texto(self.pasta / "relatorio.docx")[0])
        (self.pasta / "contrato.pdf").write_bytes(PDF)
        texto, info = documentos.extrair_texto(self.pasta / "contrato.pdf")
        self.assertIn("Contrato de aluguel 2026", texto)
        self.assertEqual(info["paginas"], 1)
        with self.assertRaises(ValueError):
            documentos.extrair_texto(self.pasta / "x.exe")
        (self.pasta / "quebrado.pdf").write_bytes(b"%PDF-1.4 lixo")
        with self.assertRaises(ValueError):                           # erro claro, nao PdfReadError cru
            documentos.extrair_texto(self.pasta / "quebrado.pdf")

    def test_achar(self):
        baixados, docs = self.pasta / "Downloads", self.pasta / "Documentos"
        (docs / "trabalho").mkdir(parents=True)
        baixados.mkdir()
        velho = baixados / "antigo.pdf"
        velho.write_bytes(PDF)
        novo = baixados / "boleto.pdf"
        novo.write_bytes(PDF)
        os.utime(velho, (time.time() - 3600, time.time() - 3600))
        (docs / "trabalho" / "Relatorio de Vendas.docx").write_bytes(b"x")
        pastas = {"downloads": baixados, "documentos": docs, "area de trabalho": self.pasta / "Desktop"}
        with mock.patch.object(documentos, "pasta_conhecida", lambda n: pastas[n]):
            self.assertEqual(documentos.achar("resuma o último PDF baixado"), novo)
            self.assertEqual(documentos.achar("resuma o relatório de vendas").name, "Relatorio de Vendas.docx")
            self.assertEqual(documentos.achar("resuma o documento aberto", "boleto.pdf - Adobe Acrobat"), novo)
            self.assertIsNone(documentos.achar("resuma o documento xyzzy"))
            self.assertEqual(documentos.achar(f"resuma {novo}"), novo)

    def test_resumo_longo_avisa_que_e_so_o_comeco(self):
        t = self.pasta / "longo.md"
        t.write_text("palavra " * 3000, encoding="utf-8")
        pedidos = []
        r = documentos.resumir(t, lambda m: pedidos.append(m) or "É um texto repetitivo.")
        self.assertTrue(r["cortado"])
        self.assertLessEqual(len(pedidos[0][1]["content"]), documentos.LIMITE_CHARS + 200)
        self.assertIn("Resumi só o começo", documentos.fala(r))


class _Seg:
    def __init__(self, a, b, t):
        self.start, self.end, self.text = a, b, t


class TestMidia(unittest.TestCase):
    def test_transcreve_por_trecho_e_grava(self):
        trava = threading.Lock()
        travado_durante = []

        def gerador():
            for s in (_Seg(0, 4, " Bom dia a todos."), _Seg(4, 9, " Vamos falar do orçamento.")):
                travado_durante.append(trava.locked())
                yield s
        modelo = SimpleNamespace(transcribe=lambda caminho, **kw: (gerador(), SimpleNamespace(duration=9.5)))
        tr = SimpleNamespace(_modelo=modelo, _trava=trava)
        audio = Path(tempfile.mkdtemp()) / "reuniao.mp3"
        audio.write_bytes(b"x")
        saida = Path(tempfile.mkdtemp())
        r = midia_arquivos.analisar(audio, tr, saida)
        self.assertEqual(travado_durante, [True, True])              # trava so durante cada trecho
        self.assertFalse(trava.locked())
        texto = Path(r["saida"]).read_text(encoding="utf-8")
        self.assertIn("[0:04] Vamos falar do orçamento.", texto)
        fala = midia_arquivos.fala(r)
        self.assertIn("Áudio reuniao.mp3, 0:09 de duração", fala)
        self.assertIn("Começa com: Bom dia a todos.", fala)


if __name__ == "__main__":
    unittest.main()
